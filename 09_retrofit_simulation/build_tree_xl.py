"""Excel workbook of the decision tree's training beams (70%) and test beams (30%).
Same split as retrofit_tree.py (70/30, stratified, random_state 7). The tree's rules are written as Excel formulas,
so Excel itself predicts the scheme of every beam and counts how many it gets right."""
import csv, numpy as np
from sklearn.tree import DecisionTreeClassifier
from sklearn.model_selection import train_test_split
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

R = list(csv.DictReader(open('retrofit_sim.csv')))
F = ['stiffness_left_lower', 'crack_ratio', 'strength_ratio', 'deflection_ratio', 'live_load_factor', 'span_depth', 'pt', 'utilisation', 'dead_fraction']
merge = {'Concrete jacket': 'Steel plate or jacket', 'Steel plate': 'Steel plate or jacket'}
X = np.array([[float(r[k]) for k in F] for r in R]); y = np.array([merge.get(r['best'], r['best']) for r in R])
idx = np.arange(len(R))
Xtr, Xte, ytr, yte, itr, ite = train_test_split(X, y, idx, test_size=0.3, random_state=7, stratify=y)
tree = DecisionTreeClassifier(max_depth=3, min_samples_leaf=20, random_state=7, class_weight='balanced').fit(Xtr, ytr)
T = tree.tree_
thr = {}
def walk(n, path):
    if T.children_left[n] == -1: return
    thr[path] = (F[T.feature[n]], float(T.threshold[n]))
    walk(T.children_left[n], path + 'L'); walk(T.children_right[n], path + 'R')
walk(0, '')
PATHS = sorted(thr, key=lambda k: (len(k), k))
def leaf(path):
    n = 0
    for c in path: n = T.children_left[n] if c == 'L' else T.children_right[n]
    return tree.classes_[int(np.argmax(T.value[n]))]
pred_sk = tree.predict(X)

AR = 'Arial'
BOLD = Font(name=AR, bold=True); BLUE = Font(name=AR, color='0000FF'); BLACK = Font(name=AR, color='000000')
GREEN = Font(name=AR, color='008000'); ITAL = Font(name=AR, italic=True, color='555555'); TITLE = Font(name=AR, bold=True, size=13)
KEY = PatternFill('solid', fgColor='FFFF00'); HEAD = PatternFill('solid', fgColor='E7E6E6')
thin = Side(style='thin', color='BFBFBF')
wb = Workbook()

# ---------------- Read me
rm = wb.active; rm.title = 'Read me'
rows = [
    ('What this workbook is', 'The %d beams of the retrofit simulation (slide 11), split as the decision tree saw them: %d training beams (70%%) and %d test beams (30%%).' % (len(R), len(itr), len(ite)) + ' The tree learned its rules from the training beams only. The test beams were never shown to it; its predictions for them are compared with the answer the IS 456 checks gave.'),
    ('How the split was made', 'Random, stratified by scheme (each scheme has the same share in both sets), random seed 7 (Python: scikit-learn train_test_split, test_size 0.3). Same split as 09_retrofit_simulation/retrofit_tree.py.'),
    ('Correct answer (column "Scheme by IS 456 checks")', 'From retrofit_sim.py: the least intrusive scheme that makes the beam comply with IS 456 cl 38.1 (flexure), cl 40 (shear), cl 23.2 (deflection) and Annex F (crack width). Order: monitor, seal cracks, polymer laminate, steel plate, concrete jacket. Steel plate and concrete jacket are one class for the tree ("Steel plate or jacket", only %d of %d beams).' % (int((y == 'Steel plate or jacket').sum()), len(R))),
    ('Tree prediction (column "Tree prediction")', 'An Excel formula: the tree\'s split values (sheet Tree rules) written as nested IF statements on the beam\'s own numbers. Excel predicts every beam itself; nothing is pasted from Python.'),
    ('Sheet: Training beams', '%d beams the tree learned from.' % len(itr)),
    ('Sheet: Test beams', '%d beams the tree never saw. Accuracy on these is the number on slide 11.' % len(ite)),
    ('Sheet: Tree rules', 'The %d split values found by the tree (green: carried from Python), and the tree written as rules.' % len(thr)),
    ('Sheet: Results', 'Accuracy, confusion tables and catch rate of plate or jacket beams, all computed by Excel formulas from the two beam sheets.'),
    ('Inputs the tree could use', 'Stiffness left (twin, lower 95%% bound), crack width / IS 456 limit, moment demand / capacity, deflection / span/250 limit, live load factor (change of use), span / effective depth, steel ratio, design utilisation, dead load fraction. It chose to split on %d of them.' % len({v[0] for v in thr.values()})),
    ('Colour code', 'Blue text = data from the simulation (do not edit). Black text = Excel formulas. Green text = split values carried from the Python tree. Yellow cells = headline results.'),
    ('Source files', '09_retrofit_simulation: retrofit_sim.py, retrofit_sim.csv, retrofit_tree.py, retrofit_tree.txt, build_tree_xl.py (this workbook).'),
]
rm['A1'] = 'Decision tree: training beams and test beams'; rm['A1'].font = TITLE
for i, (a, b) in enumerate(rows):
    rm.cell(3 + i, 1, a).font = BOLD; c = rm.cell(3 + i, 2, b); c.font = BLACK; c.alignment = Alignment(wrap_text=True, vertical='top')
    rm.cell(3 + i, 1).alignment = Alignment(wrap_text=True, vertical='top')
rm.column_dimensions['A'].width = 34; rm.column_dimensions['B'].width = 120

# ---------------- Tree rules (thresholds as cells)
tr = wb.create_sheet('Tree rules')
tr['A1'] = 'Split values found by the decision tree (depth 3, at least 20 beams per end point, classes weighted)'; tr['A1'].font = BOLD
hdr = ['Cell name', 'Quantity', 'Split value', 'Where in the tree']
for j, h in enumerate(hdr):
    c = tr.cell(3, 1 + j, h); c.font = BOLD; c.fill = HEAD
names = {'strength_ratio': 'Moment demand / capacity', 'crack_ratio': 'Crack width / IS 456 limit', 'deflection_ratio': 'Deflection / span/250 limit',
         'stiffness_left_lower': 'Stiffness left (twin, lower bound)', 'span_depth': 'Span / effective depth', 'live_load_factor': 'Live load factor',
         'pt': 'Steel ratio (%)', 'utilisation': 'Design utilisation', 'dead_fraction': 'Dead load fraction'}
def where_txt(path):
    if not path: return 'First split'
    parts = []
    n = 0
    for c in path:
        f, v = F[T.feature[n]], T.threshold[n]
        parts.append('%s %s %.2f' % (names[f], '<=' if c == 'L' else '>', v))
        n = T.children_left[n] if c == 'L' else T.children_right[n]
    return 'After: ' + '; '.join(parts)
TCELL = {}; SNAME = {}
for i_, k in enumerate(PATHS):
    r = 4 + i_; f, v = thr[k]; SNAME[k] = 'S%d' % (i_ + 1)
    tr.cell(r, 1, SNAME[k]).font = BOLD; tr.cell(r, 2, names[f]).font = BLACK
    c = tr.cell(r, 3, v); c.font = GREEN; c.number_format = '0.0000'
    tr.cell(r, 4, where_txt(k)).font = BLACK
    TCELL[k] = "'Tree rules'!$C$%d" % r
rules = []
def rule_lines(path, depth):
    if path not in thr:
        rules.append('    ' * depth + '-> ' + leaf(path)); return
    f, v = thr[path]
    rules.append('    ' * depth + 'If %s <= %s (about %.2f):' % (names[f], SNAME[path], v)); rule_lines(path + 'L', depth + 1)
    rules.append('    ' * depth + 'If %s > %s:' % (names[f], SNAME[path])); rule_lines(path + 'R', depth + 1)
rule_lines('', 0)
R0 = 5 + len(PATHS)
tr.cell(R0, 1, 'The tree as rules').font = BOLD
for i, s in enumerate(rules): tr.cell(R0 + 1 + i, 1, s).font = BLACK
tr.cell(R0 + 2 + len(rules), 1, 'Split values are shown rounded; the formulas use the full values in column C.').font = ITAL
for col, w in zip('ABCD', (12, 36, 14, 46)): tr.column_dimensions[col].width = w

# ---------------- beam sheets
COLS = [('Beam no.', 'beam', '0', 9), ('Span (m)', 'span_m', '0.00', 9), ('Width b (mm)', 'b', '0', 10), ('Depth D (mm)', 'D', '0', 10),
        ('Concrete grade fck (MPa)', 'fck', '0', 12), ('Steel ratio (%)', 'pt', '0.00', 10), ('Exposure (IS 456 Table 3)', 'exposure', None, 28),
        ('Crack limit (mm)', 'crack_limit_mm', '0.0', 9), ('Live load factor (change of use)', 'live_load_factor', '0.0', 13),
        ('Design utilisation', 'utilisation', '0.000', 11), ('Dead load fraction', 'dead_fraction', '0.000', 11),
        ('Span / effective depth', 'span_depth', '0.00', 11), ('Stiffness left (twin, lower bound)', 'stiffness_left_lower', '0.000', 13),
        ('Crack width / limit', 'crack_ratio', '0.000', 11), ('Moment demand / capacity', 'strength_ratio', '0.000', 12),
        ('Deflection / span/250 limit', 'deflection_ratio', '0.000', 12), ('Scheme by IS 456 checks', 'best', None, 18),
        ('Class used by the tree', None, None, 20), ('Tree prediction', None, None, 20), ('Correct?', None, None, 9)]
col = {h: L(j + 1) for j, (h, *_ ) in enumerate(COLS)}
num = {k for _, k, fmt, _ in COLS if k and fmt}

def beam_sheet(name, ids, note):
    ws = wb.create_sheet(name)
    ws.cell(1, 1, note).font = BOLD
    for j, (h, k, fmt, w) in enumerate(COLS):
        c = ws.cell(3, j + 1, h); c.font = BOLD; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[L(j + 1)].width = w
    ws.row_dimensions[3].height = 48
    for i, b in enumerate(ids):
        r = 4 + i; row = R[b]
        for j, (h, k, fmt, w) in enumerate(COLS):
            if k is None: continue
            v = row[k]
            v = (int(float(v)) if fmt == '0' else float(v)) if k in num else v
            c = ws.cell(r, j + 1, v); c.font = BLUE
            if fmt: c.number_format = fmt
        S, CR, SR, DR = (col['Stiffness left (twin, lower bound)'] + str(r), col['Crack width / limit'] + str(r),
                         col['Moment demand / capacity'] + str(r), col['Deflection / span/250 limit'] + str(r))
        SD = col['Span / effective depth'] + str(r); BEST = col['Scheme by IS 456 checks'] + str(r)
        ws[col['Class used by the tree'] + str(r)] = '=IF(OR(%s="Steel plate",%s="Concrete jacket"),"Steel plate or jacket",%s)' % (BEST, BEST, BEST)
        cellref = {f_: col[h_] + str(r) for f_, h_ in (('strength_ratio', 'Moment demand / capacity'), ('crack_ratio', 'Crack width / limit'),
                   ('deflection_ratio', 'Deflection / span/250 limit'), ('stiffness_left_lower', 'Stiffness left (twin, lower bound)'),
                   ('span_depth', 'Span / effective depth'), ('live_load_factor', 'Live load factor (change of use)'), ('pt', 'Steel ratio (%)'),
                   ('utilisation', 'Design utilisation'), ('dead_fraction', 'Dead load fraction'))}
        def fx(path):
            if path not in thr: return '"%s"' % leaf(path)
            return 'IF(%s<=%s,%s,%s)' % (cellref[thr[path][0]], TCELL[path], fx(path + 'L'), fx(path + 'R'))
        ws[col['Tree prediction'] + str(r)] = '=' + fx('')
        ws[col['Correct?'] + str(r)] = '=IF(%s=%s,"Yes","No")' % (col['Tree prediction'] + str(r), col['Class used by the tree'] + str(r))
        for h in ('Class used by the tree', 'Tree prediction', 'Correct?'): ws[col[h] + str(r)].font = BLACK
    ws.freeze_panes = 'B4'
    ws.auto_filter.ref = 'A3:%s%d' % (L(len(COLS)), 3 + len(ids))
    return ws, 4, 3 + len(ids)

order_tr = sorted(itr.tolist()); order_te = sorted(ite.tolist())
wtr, a0, a1 = beam_sheet('Training beams', order_tr, 'Training beams: the %d beams the tree learned its rules from' % len(itr))
wte, b0, b1 = beam_sheet('Test beams', order_te, 'Test beams: the %d beams the tree never saw while learning' % len(ite))

# ---------------- Results
rs = wb.create_sheet('Results')
rs['A1'] = 'How well the tree does (all values computed by Excel from the two beam sheets)'; rs['A1'].font = BOLD
CL = ['Monitor', 'Seal cracks', 'Polymer laminate', 'Steel plate or jacket']
def block(r0, sheet, lo, hi, title):
    rng = lambda h: "'%s'!$%s$%d:$%s$%d" % (sheet, col[h], lo, col[h], hi)
    rs.cell(r0, 1, title).font = BOLD
    rs.cell(r0 + 1, 1, 'Beams').font = BLACK; c = rs.cell(r0 + 1, 2, '=COUNTA(%s)' % rng('Beam no.')); c.font = BLACK
    rs.cell(r0 + 2, 1, 'Correct').font = BLACK; c = rs.cell(r0 + 2, 2, '=COUNTIF(%s,"Yes")' % rng('Correct?')); c.font = BLACK
    rs.cell(r0 + 3, 1, 'Accuracy').font = BOLD; c = rs.cell(r0 + 3, 2, '=B%d/B%d' % (r0 + 2, r0 + 1)); c.font = BLACK; c.number_format = '0.0%'; c.fill = KEY
    rs.cell(r0 + 5, 1, 'Confusion table: rows = scheme by IS 456 checks, columns = tree prediction').font = ITAL
    for j, cl in enumerate(CL):
        c = rs.cell(r0 + 6, 2 + j, cl); c.font = BOLD; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, horizontal='center')
    c = rs.cell(r0 + 6, 6, 'Total'); c.font = BOLD; c.fill = HEAD
    c = rs.cell(r0 + 6, 7, 'Caught (share of this scheme predicted correctly)'); c.font = BOLD; c.fill = HEAD; c.alignment = Alignment(wrap_text=True)
    for i, ci in enumerate(CL):
        rr = r0 + 7 + i; rs.cell(rr, 1, ci).font = BOLD
        for j, cj in enumerate(CL):
            c = rs.cell(rr, 2 + j, '=COUNTIFS(%s,"%s",%s,"%s")' % (rng('Class used by the tree'), ci, rng('Tree prediction'), cj)); c.font = BLACK
        c = rs.cell(rr, 6, '=SUM(B%d:E%d)' % (rr, rr)); c.font = BLACK
        c = rs.cell(rr, 7, '=IF(F{0}=0,0,{1}{0}/F{0})'.format(rr, L(2 + i))); c.font = BLACK; c.number_format = '0.0%'
    rs.cell(r0 + 11, 1, 'Balanced accuracy (mean of the four catch rates)').font = BOLD
    c = rs.cell(r0 + 11, 7, '=AVERAGE(G%d:G%d)' % (r0 + 7, r0 + 10)); c.font = BLACK; c.number_format = '0.0%'; c.fill = KEY
    rs.cell(r0 + 12, 1, 'Plate or jacket beams caught').font = BOLD
    c = rs.cell(r0 + 12, 7, '=E%d&" of "&F%d' % (r0 + 10, r0 + 10)); c.font = BLACK; c.fill = KEY
    return r0 + 14
nx = block(3, 'Test beams', b0, b1, 'Test beams (%d, never seen by the tree): the result on slide 11' % len(ite))
nx = block(nx, 'Training beams', a0, a1, 'Training beams (%d, used to learn the rules)' % len(itr))
rs.cell(nx, 1, 'Test and training accuracy are close, so the tree has not memorised the training beams.').font = ITAL
rs.cell(nx + 1, 1, 'Most errors are on the safe side: polymer-laminate beams sent to plate or jacket (a stronger scheme than needed).').font = ITAL
rs.column_dimensions['A'].width = 52
for c_ in 'BCDEF': rs.column_dimensions[c_].width = 15
rs.column_dimensions['G'].width = 24

wb.move_sheet('Results', offset=-3)
wb.save('StrucIQ_decision_tree_beams.xlsx')

# python-side check of what Excel should get
pt = np.array(pred_sk)
print('python test acc', (pt[ite] == y[ite]).mean(), 'train acc', (pt[itr] == y[itr]).mean())
np.save('tree_pred_sk.npy', pt)
